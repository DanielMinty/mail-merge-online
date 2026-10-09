"""Spreadsheet loading, plain-text mail merge, and durable send history."""
from __future__ import annotations

import csv
import hashlib
import json
import re
import sqlite3
import unicodedata
from dataclasses import asdict, dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit


class MergeError(ValueError):
    pass


def column_key(value: str) -> str:
    return re.sub(r"[^\w]+", "_", unicodedata.normalize("NFKC", value).strip().lower()).strip("_")


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@dataclass(frozen=True)
class Sheet:
    headers: tuple[str, ...]
    rows: tuple[tuple[int, dict[str, str]], ...]


def read_sheet(path: str | Path, sheet_name: str = "") -> Sheet:
    path = Path(path)
    if path.suffix.lower() in (".csv", ".tsv"):
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
            records = list(csv.reader(stream, dialect))
    elif path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise MergeError("Excel support is missing. Run the included launcher to install it.") from exc
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet_name:
                if sheet_name not in workbook.sheetnames:
                    raise MergeError(f"No worksheet named {sheet_name!r}. Available: {', '.join(workbook.sheetnames)}")
                worksheet = workbook[sheet_name]
            else:
                worksheet = next((s for s in workbook if s.sheet_state == "visible"), workbook.worksheets[0])
            records = []
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type in ("f", "e"):
                        raise MergeError(f"{worksheet.title}!{cell.coordinate} contains a formula or Excel error. Paste formulas as values and fix errors first.")
                records.append([cell.value for cell in row])
        finally:
            workbook.close()
    else:
        raise MergeError("Choose an .xlsx, .csv, or .tsv file. Save old .xls workbooks as .xlsx first.")
    if not records:
        raise MergeError("The spreadsheet is empty.")
    # Excel sometimes includes formatted but unused trailing columns.
    while records[0] and not cell_text(records[0][-1]):
        records[0] = records[0][:-1]
    headers = tuple(cell_text(v) for v in records[0])
    keys = tuple(column_key(v) for v in headers)
    if not headers or any(not k for k in keys):
        raise MergeError("Every used column needs a header in the first row.")
    if len(set(keys)) != len(keys):
        raise MergeError("Two headers become the same placeholder. Use distinct column names.")
    rows = []
    for number, values in enumerate(records[1:], 2):
        if not any(cell_text(v) for v in values):
            continue
        if any(cell_text(v) for v in values[len(keys):]):
            raise MergeError(f"Row {number} has data in a column without a header.")
        padded = list(values[:len(keys)]) + [""] * max(0, len(keys) - len(values))
        rows.append((number, {k: cell_text(v) for k, v in zip(keys, padded)}))
    if not rows:
        raise MergeError("There are headers but no recipient rows.")
    return Sheet(headers, tuple(rows))


EMAIL_RE = re.compile(r"^[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9](?:[A-Za-z0-9.-]*[A-Za-z0-9])?$")
EMAIL_IN_TEXT = re.compile(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9](?:[A-Za-z0-9.-]*[A-Za-z0-9])?")


def email_address(value: str) -> str:
    value = value.strip()
    if not EMAIL_RE.fullmatch(value) or len(value) > 254:
        raise MergeError(f"Invalid email address: {value!r}. Use a plain address, without a display name.")
    local, domain = value.rsplit("@", 1)
    if local.startswith(".") or local.endswith(".") or ".." in local or len(local) > 64:
        raise MergeError(f"Invalid email address: {value!r}")
    if any(not label or len(label) > 63 or label.startswith("-") or label.endswith("-") for label in domain.split(".")):
        raise MergeError(f"Invalid email address: {value!r}")
    return local + "@" + domain.lower()


def address_list(value: str) -> tuple[str, ...]:
    if not value.strip():
        return ()
    result = {}
    for item in re.split(r"[,;]", value):
        if not item.strip():
            continue
        address = email_address(item)
        result.setdefault(address.casefold(), address)
    return tuple(result.values())


PLACEHOLDER = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def render_template(template: str, values: dict[str, str], row: int) -> str:
    def replace(match):
        key = column_key(match.group(1))
        if key not in values:
            raise MergeError(f"Unknown placeholder {match.group(0)!r}. Use a spreadsheet column header.")
        if not values[key]:
            raise MergeError(f"Row {row}: {match.group(0)} is blank. Fill that cell before sending.")
        return values[key]
    # A single pass prevents cell contents from becoming a second template.
    return PLACEHOLDER.sub(replace, template)


@dataclass(frozen=True)
class Message:
    row: int
    to: str
    cc: tuple[str, ...]
    subject: str
    body: str


def merge(sheet: Sheet, email_column: str, cc_column: str, global_cc: str,
          subject: str, body: str) -> tuple[Message, ...]:
    email_key = column_key(email_column)
    cc_key = column_key(cc_column)
    keys = {column_key(h) for h in sheet.headers}
    if email_key not in keys or (cc_key and cc_key not in keys):
        raise MergeError("Choose the Email and optional CC columns from the spreadsheet.")
    if not subject.strip() or not body.strip():
        raise MergeError("Enter a subject and message body.")
    fixed_cc = address_list(global_cc)
    seen = {}
    result = []
    for number, values in sheet.rows:
        try:
            recipient = email_address(values[email_key])
            normalized = recipient.casefold()
            if normalized in seen:
                raise MergeError(f"Duplicate recipient on rows {seen[normalized]} and {number}: {recipient}")
            seen[normalized] = number
            cc = dict.fromkeys(a.casefold() for a in (*fixed_cc, *address_list(values.get(cc_key, ""))))
            cc.pop(normalized, None)
            rendered_subject = render_template(subject, values, number).strip()
            rendered_body = render_template(body, values, number)
            if any(ord(c) < 32 for c in rendered_subject):
                raise MergeError("The subject contains a line break or control character.")
            result.append(Message(number, recipient, tuple(cc), rendered_subject, rendered_body))
        except MergeError as exc:
            raise MergeError(f"Row {number}: {exc}") from exc
    return tuple(result)


def validate_url(url: str) -> str:
    url = url.strip()
    parts = urlsplit(url)
    if parts.scheme != "https" or not parts.hostname or parts.username or parts.password:
        raise MergeError("Enter the HTTPS address of your organization's Outlook Web App.")
    return url


def profile_key(url: str, mailbox: str) -> str:
    return hashlib.sha256((url + "\n" + mailbox.casefold()).encode()).hexdigest()[:20]


def campaign_key(url: str, mailbox: str, messages: tuple[Message, ...]) -> str:
    # Reordering spreadsheet rows does not reset the history. Editing content does.
    content = [{k: v for k, v in asdict(m).items() if k != "row"} for m in messages]
    content.sort(key=lambda m: m["to"].casefold())
    payload = json.dumps([url, mailbox.casefold(), content], sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(payload.encode()).hexdigest()


class Ledger:
    """Write 'attempting' durably BEFORE clicking Send; never auto-retry it."""
    def __init__(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.connection = sqlite3.connect(path, timeout=5)
        self.connection.execute("PRAGMA synchronous=FULL")
        self.connection.execute("CREATE TABLE IF NOT EXISTS messages (campaign TEXT, recipient TEXT, row_number INTEGER, subject TEXT, status TEXT, detail TEXT, updated TEXT, PRIMARY KEY (campaign,recipient))")
        self.connection.commit()
        try:
            path.chmod(0o600)
        except OSError:
            pass

    def close(self):
        self.connection.close()

    def status(self, campaign: str, recipient: str) -> str:
        row = self.connection.execute("SELECT status FROM messages WHERE campaign=? AND recipient=?", (campaign, recipient.casefold())).fetchone()
        return row[0] if row else "pending"

    def record(self, campaign: str, message: Message, status: str, detail: str = ""):
        if status not in ("attempting", "submitted", "uncertain", "pending"):
            raise MergeError("Invalid history status.")
        with self.connection:
            self.connection.execute("INSERT INTO messages VALUES (?,?,?,?,?,?,?) ON CONFLICT(campaign,recipient) DO UPDATE SET row_number=excluded.row_number,subject=excluded.subject,status=excluded.status,detail=excluded.detail,updated=excluded.updated",
                                    (campaign, message.to.casefold(), message.row, message.subject, status, detail, datetime.now(timezone.utc).isoformat()))

    def list(self, campaign: str):
        return self.connection.execute("SELECT recipient,row_number,subject,status,detail,updated FROM messages WHERE campaign=? ORDER BY row_number", (campaign,)).fetchall()

    def resolve(self, campaign: str, recipient: str, submitted: bool):
        with self.connection:
            self.connection.execute("UPDATE messages SET status=?, detail=?, updated=? WHERE campaign=? AND recipient=? AND status IN ('attempting','uncertain')",
                                    ("submitted" if submitted else "pending", "Manually checked in Sent Items", datetime.now(timezone.utc).isoformat(), campaign, recipient.casefold()))


def csv_safe(value) -> str:
    value = str(value)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value
